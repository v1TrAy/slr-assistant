
import io
import sqlite3
import pytest
import slr_core
from slr_core.dedup import title_key, DuplicateIndex
from slr_core.repository import RecordRepository


def _rec(title, **kw):
    base = {"title": title, "authors": "A", "year": "2020", "journal": "J",
            "doi": "", "database": "Test"}
    base.update(kw)
    return base


@pytest.fixture
def repo():
    r = RecordRepository(":memory:")
    yield r
    r.close()


# --- C2: dedup normalization ------------------------------------------

def test_title_key_keeps_non_latin_scripts():
    assert title_key("机器学习导论") != ""
    assert title_key("机器学习导论") != title_key("深度学习综述")


def test_empty_titles_never_match():
    idx = DuplicateIndex().build([(1, "!!! ..."), (2, "机器学习")])
    rid, score = idx.query("   ")
    assert rid is None and score == 0.0


def test_two_distinct_cjk_titles_are_not_duplicates():
    idx = DuplicateIndex().build([(1, "深度学习在医学影像中的应用")])
    rid, _ = idx.query("自然语言处理综述")
    assert rid is None


def test_accented_and_decomposed_forms_match():
    assert title_key("Café") == title_key("Cafe\u0301")


# --- M5: blocking recall ------------------------------------------------

def test_suffix_key_catches_leading_article():
    idx = DuplicateIndex().build([(1, "Effect of X on Y")])
    rid, score = idx.query("The Effect of X on Y", threshold=0.85)
    assert rid == 1 and score >= 0.85


def test_gram_index_catches_head_and_tail_edits():
    idx = DuplicateIndex().build(
        [(1, "New deep learning models for automated medical imaging analysis of chest radiographs")])
    rid, score = idx.query(
        "Old deep learning models for automated medical imaging analysis of chest radiography")
    assert rid == 1  # prefix AND suffix differ; the shared middle still links them


def test_insertion_is_found_without_exhaustive_search():
    # The old fixed-window blocking missed this pair unless exhaustive=True.
    idx = DuplicateIndex().build([(1, "Machine learning approaches")])
    blocked_rid, score = idx.query("A machine learning approach")
    assert blocked_rid == 1 and score >= 0.9
    rid, _ = idx.query("A machine learning approach", exhaustive=True)
    assert rid == 1


# --- repository ------------------------------------------------------------

def test_bulk_negative_refs(repo):
    ids = repo.insert_records_bulk([
        (_rec("A"), "No", None),
        (_rec("B"), "Yes", -1),  # duplicate of the first record in this batch
        (_rec("C"), "No", None),
    ])
    rows = {r["id"]: r for r in repo.fetch_all_for_export()}
    assert rows[ids[1]]["duplicate_of"] == ids[0]


def test_delete_unflags_children(repo):
    ids = repo.insert_records_bulk([(_rec("A"), "No", None), (_rec("A2"), "Yes", -1)])
    repo.delete_records([ids[0]])
    row = repo.fetch_all_for_export()[0]
    assert row["is_duplicate"] == "No" and row["duplicate_of"] is None


def test_restore_snapshot_rejects_unknown_columns(repo):
    repo.insert_record(_rec("A"))
    with pytest.raises(ValueError):
        repo.restore_snapshot({"cols": ["id; DROP TABLE records"], "rows": []})
    assert repo.prisma_counts()["Records identified through database searching"] == 1


def test_restore_snapshot_roundtrip(repo):
    repo.insert_record(_rec("A"))
    snap = repo.snapshot()
    repo.insert_record(_rec("B"))
    repo.restore_snapshot(snap)
    assert set(repo.all_titles_map().values()) == {"A"}


def test_update_decision_whitelists_fields(repo):
    rid = repo.insert_record(_rec("A"))
    repo.update_decision(rid, "ta_decision", "Include")
    with pytest.raises(ValueError):
        repo.update_decision(rid, "title; DROP TABLE records", "x")


def test_rejects_non_sqlite_file(tmp_path):
    bad = tmp_path / "project.xlsx"
    bad.write_bytes(b"PK\x03\x04 definitely not a database")
    with pytest.raises(ValueError):
        RecordRepository(str(bad))


# --- staging (M2) -------------------------------------------------------------

class FakeStore:
    def __init__(self):
        self.items = None

    def insert_records_bulk(self, items):
        self.items = items
        return list(range(1, len(items) + 1))


def test_import_batch_flags_within_batch_duplicates():
    batch = slr_core.ImportBatch([], threshold=0.9)
    batch.add_records([_rec("Alpha"), _rec("Alpha"), _rec("Beta")])
    assert batch.records[1].is_dup == "Yes"
    assert batch.records[1].dup_ref == -1


def test_import_batch_skipped_target_drops_dup_flag():
    batch = slr_core.ImportBatch([], threshold=0.9)
    batch.add_records([_rec("Alpha"), _rec("Alpha")])
    batch.records[0].skipped = True  # user unchecks the ORIGINAL
    fake = FakeStore()
    n_total, n_dup = batch.commit(fake)
    # The unchecked original is NOT imported; the checked duplicate IS
    # imported, but as a unique record (its duplicate target is gone).
    assert n_total == 1
    assert n_dup == 0
    assert fake.items[0][1] == "No"     # imported as unique
    assert fake.items[0][2] is None     # no dangling batch reference


def test_import_batch_skip_the_duplicate_keeps_original():
    batch = slr_core.ImportBatch([], threshold=0.9)
    batch.add_records([_rec("Alpha"), _rec("Alpha")])
    batch.records[1].skipped = True  # user unchecks the DUPLICATE
    fake = FakeStore()
    n_total, n_dup = batch.commit(fake)
    assert n_total == 1 and n_dup == 0
    assert fake.items[0][2] is None


def test_import_batch_renumbers_refs_across_skips():
    batch = slr_core.ImportBatch([], threshold=0.9)
    batch.add_records([_rec("Alpha"), _rec("Beta"), _rec("Beta")])
    # record 2 is flagged as a duplicate of record 1 (batch ref -2).
    batch.records[0].skipped = True  # skip the unrelated "Alpha"
    fake = FakeStore()
    n_total, n_dup = batch.commit(fake)
    # kept = [Beta, Beta-dup]; the duplicate's reference must be renumbered
    # from -2 to -1 so it still points at "Beta" within the committed items.
    assert n_total == 2 and n_dup == 1
    assert fake.items[0][2] is None
    assert fake.items[1][2] == -1
    assert fake.items[1][1] == "Yes"


# --- parsers (M7, L1, L2) ------------------------------------------------------

def test_csv_cp1252_fallback():
    data = "Title,Authors\nCafé Study,René Descartes\n".encode("cp1252")
    recs = slr_core.parse_csv_bytes(io.BytesIO(data), "Test")
    assert recs[0]["authors"] == "René Descartes"


def test_ris_flushes_final_record_without_er():
    ris = "TY - JOUR\nTI - Last record\nER - \nTY - JOUR\nTI - Truncated record\n"
    recs = slr_core.parse_ris_bytes(io.StringIO(ris), "Test")
    assert [r["title"] for r in recs] == ["Last record", "Truncated record"]


def test_xlsx_duplicate_headers_do_not_collapse():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Title", "Title"])
    ws.append(["Real title", "Shadow title"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    recs = slr_core.parse_xlsx_bytes(buf, "Test")
    assert len(recs) == 1 and recs[0]["title"] == "Real title"


# --- excel_io (L4) ---------------------------------------------------------------

def test_resolve_columns_prevents_cross_field_collision():
    import openpyxl
    from slr_core.excel_io import _resolve_columns
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Source ID", "Notes"])
    cols = _resolve_columns(ws, [("id", "ID"), ("database", "Database"),
                                 ("reason", "Exclusion Reason")])
    # "Source ID" word-matches both the `id` and `database` aliases; the
    # two fields must never end up writing into the same column.
    assert cols["id"] != cols["database"]


# --- DOI normalization ---------------------------------------------------

def test_normalize_doi_forms():
    from slr_core.dedup import normalize_doi
    assert normalize_doi("10.1109/ICCV.2019.00001") == "10.1109/iccv.2019.00001"
    assert normalize_doi("https://doi.org/10.1109/ICCV.2019.00001") == "10.1109/iccv.2019.00001"
    assert normalize_doi("http://dx.doi.org/10.1111/1467-9566.12345") == "10.1111/1467-9566.12345"
    assert normalize_doi("doi:10.1007/s11276-019-0203-4") == "10.1007/s11276-019-0203-4"
    assert normalize_doi("10.1016/j.cell.2016.07.054.") == "10.1016/j.cell.2016.07.054"
    assert normalize_doi("https://link.springer.com/article/10.1007/s00348-019-0281-z") \
        == "10.1007/s00348-019-0281-z"
    assert normalize_doi(" 10.1145/3290605.3300846 ") == "10.1145/3290605.3300846"


def test_normalize_doi_rejects_non_dois():
    from slr_core.dedup import normalize_doi
    assert normalize_doi("https://ieeexplore.ieee.org/document/8765432") is None
    assert normalize_doi("https://pubmed.ncbi.nlm.nih.gov/32191259/") is None
    assert normalize_doi("") is None
    assert normalize_doi(None) is None
    assert normalize_doi("no doi here") is None


# --- DOI-first deduplication ----------------------------------------------

def test_doi_match_beats_completely_different_titles():
    batch = slr_core.ImportBatch(
        existing_titles=[(7, "Completely Different Title")],
        existing_dois=[(7, "10.1109/iccv.2019.00001")])
    batch.add_records([_rec("Totally Unrelated Title",
                            doi="https://doi.org/10.1109/ICCV.2019.00001")])
    r = batch.records[0]
    assert r.is_dup == "Yes" and r.dup_ref == 7
    assert r.match_source == "doi" and r.score == 1.0
    assert r.matched_title == "Completely Different Title"


def test_doi_within_batch_three_copies():
    batch = slr_core.ImportBatch([], threshold=0.9)
    batch.add_records([
        _rec("Alpha", doi="10.1000/a"),
        _rec("Alpha version two", doi="10.1000/a"),
        _rec("Third variant", doi="10.1000/a"),
    ])
    assert batch.records[1].is_dup == "Yes" and batch.records[1].dup_ref == -1
    assert batch.records[2].is_dup == "Yes" and batch.records[2].dup_ref == -1
    assert all(r.match_source == "doi" for r in batch.records[1:])


def test_doi_veto_keeps_same_title_different_doi_unique():
    batch = slr_core.ImportBatch(
        existing_titles=[(1, "Machine Learning for X")],
        existing_dois=[(1, "10.1000/aaa")])
    batch.add_records([_rec("Machine Learning for X", doi="10.1000/bbb")])
    r = batch.records[0]
    assert r.is_dup == "No" and r.dup_ref is None
    assert r.match_source == "doi_veto"
    assert r.matched_title == "Machine Learning for X"


def test_doi_veto_disabled_flags_title_match():
    batch = slr_core.ImportBatch(
        existing_titles=[(1, "Machine Learning for X")],
        existing_dois=[(1, "10.1000/aaa")], doi_veto=False)
    batch.add_records([_rec("Machine Learning for X", doi="10.1000/bbb")])
    assert batch.records[0].is_dup == "Yes"
    assert batch.records[0].match_source == "title"


def test_no_veto_when_candidate_has_no_doi():
    batch = slr_core.ImportBatch(existing_titles=[(1, "Machine Learning for X")])
    batch.add_records([_rec("Machine Learning for X", doi="10.1000/bbb")])
    assert batch.records[0].is_dup == "Yes"  # title match stands
    assert batch.records[0].match_source == "title"


def test_uncertain_excludes_doi_vetoed_records():
    batch = slr_core.ImportBatch(
        existing_titles=[(1, "Machine Learning for X")],
        existing_dois=[(1, "10.1000/aaa")])
    batch.add_records([_rec("Machine Learning for Xy", doi="10.1000/bbb"),  # vetoed ~97%
                       _rec("Machine Learning for Xz", doi="")])            # fuzzy dup
    titles = [t for t, _, _ in batch.uncertain]
    assert "Machine Learning for Xy" not in titles  # vetoed -> not uncertain
    assert "Machine Learning for Xz" in titles      # flagged, <100% -> uncertain


def test_commit_preserves_doi_refs_to_existing_records():
    batch = slr_core.ImportBatch([], existing_dois=[(5, "10.1000/x")])
    batch.add_records([_rec("Whatever", doi="10.1000/x")])
    fake = FakeStore()
    n_total, n_dup = batch.commit(fake)
    assert n_total == 1 and n_dup == 1
    assert fake.items[0][2] == 5  # positive ref passes through


# --- repository.existing_dois ----------------------------------------------

def test_existing_dois_remaps_duplicates_to_original(repo):
    a = repo.insert_record(_rec("Alpha", doi="10.1000/a"))
    b = repo.insert_record(_rec("Alpha", doi="https://doi.org/10.1000/a"))
    repo.set_duplicate_flag(b, "Yes", a)
    assert dict(repo.existing_dois()) == {a: "10.1000/a"}


def test_existing_dois_duplicate_with_own_doi_points_to_original(repo):
    a = repo.insert_record(_rec("Alpha", doi=""))           # original, no DOI
    b = repo.insert_record(_rec("Alpha", doi="10.1000/b"))  # dup with its own DOI
    repo.set_duplicate_flag(b, "Yes", a)
    assert dict(repo.existing_dois()) == {a: "10.1000/b"}  # remapped to original


# --- new fixes: formula-injection sanitization -----------------------------

def test_xlsx_safe_neutralizes_formula_triggers():
    from slr_core.excel_io import _xlsx_safe
    assert _xlsx_safe("=cmd|'/c calc'!A1") == "'=cmd|'/c calc'!A1"
    assert _xlsx_safe("+1+1") == "'+1+1"
    assert _xlsx_safe("-1+1") == "'-1+1"
    assert _xlsx_safe("@SUM(A1)") == "'@SUM(A1)"
    assert _xlsx_safe("Normal Title") == "Normal Title"
    assert _xlsx_safe("") == ""
    assert _xlsx_safe(None) is None


def test_export_prisma_xlsx_sanitizes_malicious_title(tmp_path, repo):
    from slr_core.excel_io import export_prisma_xlsx
    import openpyxl
    repo.insert_record(_rec('=HYPERLINK("http://evil")', authors="A"))
    out = tmp_path / "report.xlsx"
    export_prisma_xlsx(repo, str(out))
    wb = openpyxl.load_workbook(str(out))
    title_cell = wb["All_Records"].cell(row=2, column=2).value
    assert not title_cell.startswith("=")
    assert title_cell.startswith("'=")


# --- new fixes: DuplicateIndex bucket cap -----------------------------------

def test_recall_does_not_depend_on_import_order():
    # The old per-bucket cap (500) silently stopped indexing records once a
    # common prefix/suffix bucket filled up, so a duplicate of a LATE record
    # was missed.
    from slr_core.dedup import DuplicateIndex
    idx = DuplicateIndex()
    for i in range(1500):
        idx.add(i, "The effect of term{} on the outcome of study analysis".format(i))
    idx.add("late", "The effect of zebrafish stripes on the outcome of study analysis")
    rid, score = idx.query("The effect of zebrafish stripe on the outcome of study analysis")
    assert rid == "late" and score > 0.95


# --- new fixes: repository.existing_dois() (previously missing entirely) ---

def test_existing_dois_basic_roundtrip(repo):
    rid = repo.insert_record(_rec("Alpha", doi="10.1000/xyz"))
    assert dict(repo.existing_dois()) == {rid: "10.1000/xyz"}


def test_existing_dois_feeds_import_batch_directly(repo):
    repo.insert_record(_rec("Something Else", doi="10.1000/zzz"))
    batch = slr_core.ImportBatch(repo.existing_titles(), repo.existing_dois())
    batch.add_records([_rec("Completely Different", doi="10.1000/zzz")])
    assert batch.records[0].is_dup == "Yes"
    assert batch.records[0].match_source == "doi"


# --- new fixes: schema v2 (CHECK constraint + doi index) -------------------

def test_schema_v2_rejects_invalid_is_duplicate_value(repo):
    with pytest.raises(sqlite3.IntegrityError):
        repo.conn.execute(
            "INSERT INTO records(title, is_duplicate) VALUES ('X', 'Maybe')")


def test_schema_v2_has_doi_index(repo):
    rows = repo.conn.execute("SELECT name FROM sqlite_master WHERE type='index'").fetchall()
    names = {r[0] for r in rows}
    assert "idx_records_doi" in names


# --- new fixes: position staleness tracking ---------------------------------

def test_exported_ids_are_stable_after_delete(repo):
    # id_to_pos_map used to be a 1-based rank, so deleting record A silently
    # renumbered B from 2 to 1 in every later export/tracker.
    a = repo.insert_record(_rec("A"))
    b = repo.insert_record(_rec("B"))
    repo.delete_records([a])
    assert repo.id_to_pos_map()[b] == b == 2
    assert repo.positions_stale() is False


def test_restore_snapshot_never_reuses_ids(repo):
    repo.insert_record(_rec("A"))
    repo.insert_record(_rec("B"))
    snap = repo.snapshot()
    c = repo.insert_record(_rec("C"))
    repo.restore_snapshot(snap)            # undo of the insert
    d = repo.insert_record(_rec("D"))
    assert d > c                           # id 3 (already handed out) is not recycled


# --- new fixes: parser skip-count stats -------------------------------------

def test_parse_csv_bytes_reports_skipped_count_via_stats():
    data = "Title,Authors\n,Someone\nReal Title,Author\n".encode("utf-8")
    stats = {}
    recs = slr_core.parse_csv_bytes(io.BytesIO(data), "Test", stats=stats)
    assert len(recs) == 1
    assert stats["skipped"] == 1


def test_parse_ris_bytes_reports_skipped_count_via_stats():
    ris = "TY - JOUR\nAU - No Title Here\nER - \nTY - JOUR\nTI - Has A Title\nER - \n"
    stats = {}
    recs = slr_core.parse_ris_bytes(io.StringIO(ris), "Test", stats=stats)
    assert len(recs) == 1
    assert stats["skipped"] == 1
