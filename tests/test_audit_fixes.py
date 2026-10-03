# -*- coding: utf-8 -*-
"""Regression tests for the issues found in the full audit (see AUDIT.md)."""

import io
import sqlite3
import time

import pytest

import slr_core
from slr_core import parsers
from slr_core.dedup import DuplicateIndex, title_key
from slr_core.repository import RecordRepository


def _rec(title, **kw):
    base = {"title": title, "authors": "A", "year": "2020", "journal": "J",
            "doi": "", "abstract": "", "database": "Test"}
    base.update(kw)
    return base


@pytest.fixture
def repo():
    r = RecordRepository(":memory:")
    yield r
    r.close()


# ---------------------------------------------------------------- dedup ----

def test_numeric_guard_keeps_part_1_and_part_2_apart():
    idx = DuplicateIndex().build([(1, "Deep learning for medical imaging, Part 1")])
    rid, _ = idx.query("Deep learning for medical imaging, Part 2")
    assert rid is None
    # ...unless explicitly disabled
    idx2 = DuplicateIndex(numeric_guard=False).build([(1, "Deep learning for medical imaging, Part 1")])
    assert idx2.query("Deep learning for medical imaging, Part 2")[0] == 1


def test_numeric_guard_treats_persian_and_ascii_digits_alike():
    idx = DuplicateIndex().build([(1, "\u0645\u0631\u0648\u0631 \u0646\u0638\u0627\u0645\u0645\u0646\u062f \u0628\u062e\u0634 2")])
    rid, _ = idx.query("\u0645\u0631\u0648\u0631 \u0646\u0638\u0627\u0645\u0645\u0646\u062f \u0628\u062e\u0634 \u06f2")
    assert rid == 1


def test_exact_title_is_found_in_constant_time_path():
    idx = DuplicateIndex().build([(7, "Some Title: With, Punctuation!")])
    assert idx.query("some title with punctuation") == (7, 1.0)


def test_persian_title_with_typo_is_found():
    base = "\u0628\u0631\u0631\u0633\u06cc \u0645\u0631\u0648\u0631\u06cc \u0646\u0638\u0627\u0645\u0645\u0646\u062f \u06cc\u0627\u062f\u06af\u06cc\u0631\u06cc \u0645\u0627\u0634\u06cc\u0646"
    typo = base.replace("\u0645\u0627\u0634\u06cc\u0646", "\u0645\u0627\u0634\u06cc\u0646\u0647")
    idx = DuplicateIndex().build([(1, base)])
    assert idx.query(typo)[0] == 1


def test_dedup_scales_roughly_linearly():
    # 3,000 records took ~19 s with the old blocking index; keep a very
    # generous ceiling so the test is not machine-dependent.
    import random
    rng = random.Random(1)
    words = ["w%04d" % i for i in range(3000)] + ["of", "the", "and", "for", "learning"]
    idx = DuplicateIndex()
    t0 = time.perf_counter()
    for i in range(3000):
        title = " ".join(rng.choice(words) for _ in range(10))
        idx.query(title)
        idx.add(i, title)
    assert time.perf_counter() - t0 < 10


# -------------------------------------------------------------- parsers ----

def test_cp1256_arabic_csv_is_not_mojibake():
    text = "Title,Authors\n\u0645\u0631\u0627\u062c\u0639\u0629 \u0645\u0646\u0647\u062c\u064a\u0629 \u0644\u0644\u0623\u062f\u0628\u064a\u0627\u062a,\u0623\u062d\u0645\u062f\n"
    stats = {}
    recs = slr_core.parse_csv_bytes(io.BytesIO(text.encode("cp1256")), "T", stats=stats)
    assert recs[0]["title"] == "\u0645\u0631\u0627\u062c\u0639\u0629 \u0645\u0646\u0647\u062c\u064a\u0629 \u0644\u0644\u0623\u062f\u0628\u064a\u0627\u062a"
    assert any("cp1256" in g for g in stats["encoding_guesses"])


def test_utf8_persian_needs_no_guess():
    text = "Title\n\u0628\u0631\u0631\u0633\u06cc \u0646\u0638\u0627\u0645\u0645\u0646\u062f\n"
    stats = {}
    recs = slr_core.parse_csv_bytes(io.BytesIO(text.encode("utf-8")), "T", stats=stats)
    assert recs[0]["title"] == "\u0628\u0631\u0631\u0633\u06cc \u0646\u0638\u0627\u0645\u0645\u0646\u062f"
    assert "encoding_guesses" not in stats


def test_utf16_with_bom_is_decoded():
    data = "Title,Year\nUnicode Study,2021\n".encode("utf-16")
    recs = slr_core.parse_csv_bytes(io.BytesIO(data), "T")
    assert recs[0]["title"] == "Unicode Study" and recs[0]["year"] == "2021"


def test_western_cp1252_still_uses_cp1252_not_cp1256():
    data = "Title,Authors\nCaf\u00e9 Study,Ren\u00e9 Descartes\n".encode("cp1252")
    stats = {}
    recs = slr_core.parse_csv_bytes(io.BytesIO(data), "T", stats=stats)
    assert recs[0]["authors"] == "Ren\u00e9 Descartes"
    assert "cp1252" in stats["encoding_guesses"][0]


def test_semicolon_delimited_csv_is_parsed():
    data = "Title;Authors;Year\nSemi Study;Someone;2019\n".encode("utf-8")
    recs = slr_core.parse_csv_bytes(io.BytesIO(data), "T")
    assert len(recs) == 1 and recs[0]["year"] == "2019"


def test_csv_doi_beats_earlier_publisher_url():
    data = ("Title,URL,DOI\nA Paper,https://publisher.example/a/1,10.1000/real.doi\n").encode("utf-8")
    assert slr_core.parse_csv_bytes(io.BytesIO(data), "T")[0]["doi"] == "10.1000/real.doi"


def test_csv_url_kept_when_no_doi_exists():
    data = "Title,URL\nA Paper,https://publisher.example/a/1\n".encode("utf-8")
    assert slr_core.parse_csv_bytes(io.BytesIO(data), "T")[0]["doi"] == "https://publisher.example/a/1"


def test_ris_doi_beats_earlier_ur_line():
    ris = "TY  - JOUR\nTI  - X\nUR  - https://publisher.example/x\nDO  - 10.1000/abc\nER  - \n"
    assert slr_core.parse_ris_bytes(io.StringIO(ris), "T")[0]["doi"] == "10.1000/abc"


def test_abstract_is_captured_from_csv_and_ris_including_wrapped_lines():
    csv_data = 'Title,Abstract\nP,"Line one.\nLine two."\n'.encode("utf-8")
    assert slr_core.parse_csv_bytes(io.BytesIO(csv_data), "T")[0]["abstract"] == "Line one.\nLine two."
    ris = "TY  - JOUR\nTI  - P\nAB  - First part of the\nsecond part\nPY  - 2020\nER  - \n"
    rec = slr_core.parse_ris_bytes(io.StringIO(ris), "T")[0]
    assert rec["abstract"] == "First part of the second part" and rec["year"] == "2020"


def test_ris_record_without_ty_line_is_kept():
    ris = "TI  - No type line\nAU  - Someone\nER  - \n"
    recs = slr_core.parse_ris_bytes(io.StringIO(ris), "T")
    assert recs[0]["title"] == "No type line" and recs[0]["database"] == "T"


def test_year_regex_does_not_match_inside_longer_numbers():
    assert parsers._extract_year("ref 120195") == "ref 120195"
    assert parsers._extract_year("Published 2021-03") == "2021"


@pytest.mark.parametrize("name,expected", [
    ("export2026.09.18-08.33.53.csv", "IEEE Xplore"),
    ("savedrecs.ris", "Web of Science"),
    ("scopus_export.csv", "Scopus"),
    ("pubmed-covid-set.ris", "PubMed"),
    ("ieee2024.csv", "IEEE Xplore"),
    ("acme-search.csv", "acme-search"),
    ("my_notes.csv", "my_notes"),
])
def test_guess_db_name(name, expected):
    assert slr_core.guess_db_name(name) == expected


def test_real_ieee_style_export_keeps_abstract(tmp_path):
    p = tmp_path / "export2026.01.01-00.00.00.csv"
    p.write_bytes(('"Document Title",Authors,"Publication Title","Publication Year",Abstract,DOI,"PDF Link"\n'
                   '"T1","A; B","Conf",2025,"An abstract.",10.1109/X.1,https://ieeexplore.ieee.org/stamp/1\n'
                   ).encode("utf-8"))
    with open(p, "rb") as fh:
        recs = slr_core.normalize_uploaded_file(p.name, fh)
    assert recs[0]["abstract"] == "An abstract."
    assert recs[0]["doi"] == "10.1109/X.1" and recs[0]["database"] == "IEEE Xplore"


# ----------------------------------------------------------- repository ----

V2_DDL = """
CREATE TABLE records (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT, authors TEXT, year TEXT, journal TEXT, database_name TEXT, doi TEXT,
    is_duplicate TEXT NOT NULL DEFAULT 'No' CHECK (is_duplicate IN ('Yes', 'No')),
    duplicate_of INTEGER,
    ta_decision TEXT DEFAULT '', ta_reason TEXT DEFAULT '',
    ft_decision TEXT DEFAULT '', ft_reason TEXT DEFAULT '',
    study_design TEXT DEFAULT '', population TEXT DEFAULT '', variables TEXT DEFAULT '',
    findings TEXT DEFAULT '', limitations TEXT DEFAULT '', notes TEXT DEFAULT ''
);
"""


def _make_v2_db(path):
    c = sqlite3.connect(path)
    c.executescript(V2_DDL)
    c.execute("INSERT INTO records(title, ta_decision) VALUES ('Old record', 'Include')")
    c.execute("PRAGMA user_version = 2")
    c.commit()
    c.close()


def test_v2_database_migrates_to_v3_with_backup(tmp_path):
    path = str(tmp_path / "p.slrdb")
    _make_v2_db(path)
    with RecordRepository(path) as r:
        assert r.conn.execute("PRAGMA user_version").fetchone()[0] == 3
        assert r.all_titles_map() == {1: "Old record"}
        names = {row[0] for row in r.conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert {"record_abstracts", "decision_log", "meta"} <= names
    backup = sqlite3.connect(path + ".pre-v3-migration.bak")
    assert backup.execute("PRAGMA user_version").fetchone()[0] == 2
    assert backup.execute("SELECT COUNT(*) FROM records").fetchone()[0] == 1
    backup.close()


def test_failed_migration_rolls_back_completely(tmp_path, monkeypatch):
    from slr_core import repository
    path = str(tmp_path / "p.slrdb")
    _make_v2_db(path)
    monkeypatch.setitem(repository.MIGRATIONS, 3, [
        "CREATE TABLE record_abstracts (record_id INTEGER PRIMARY KEY, abstract TEXT NOT NULL)",
        "THIS IS NOT SQL",
    ])
    with pytest.raises(sqlite3.Error):
        RecordRepository(path)
    c = sqlite3.connect(path)
    assert c.execute("PRAGMA user_version").fetchone()[0] == 2           # still v2
    assert c.execute("SELECT COUNT(*) FROM records").fetchone()[0] == 1  # data intact
    assert c.execute("SELECT COUNT(*) FROM sqlite_master WHERE name='record_abstracts'").fetchone()[0] == 0
    c.close()


def test_foreign_sqlite_database_is_not_adopted(tmp_path):
    path = str(tmp_path / "browser.sqlite")
    c = sqlite3.connect(path)
    c.execute("CREATE TABLE cookies (host TEXT)")
    c.commit()
    c.close()
    with pytest.raises(ValueError):
        RecordRepository(path)
    c = sqlite3.connect(path)
    names = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")]
    c.close()
    assert names == ["cookies"]  # untouched


def test_newer_schema_is_refused(tmp_path):
    path = str(tmp_path / "future.slrdb")
    RecordRepository(path).close()
    c = sqlite3.connect(path)
    c.execute("PRAGMA user_version = 99")
    c.commit()
    c.close()
    with pytest.raises(RuntimeError):
        RecordRepository(path)


def test_abstract_roundtrip_and_undo_keeps_it(repo):
    rid = repo.insert_record(_rec("A", abstract="  The abstract text.  "))
    assert repo.get_abstract(rid) == "The abstract text."
    snap = repo.snapshot()
    repo.delete_records([rid])
    repo.restore_snapshot(snap)                 # undo of a delete
    assert repo.get_abstract(rid) == "The abstract text."


def test_bulk_insert_stores_abstracts_and_logs_the_import(repo):
    ids = repo.insert_records_bulk([(_rec("A", abstract="x"), "No", None), (_rec("B"), "No", None)])
    assert repo.get_abstract(ids[0]) == "x" and repo.get_abstract(ids[1]) == ""
    entries = [tuple(r)[3:] for r in repo.fetch_decision_log() if r["field"] == "__import__"]
    assert len(entries) == 1


def test_orphan_abstracts_are_purged_on_reopen(tmp_path):
    path = str(tmp_path / "p.slrdb")
    with RecordRepository(path) as r:
        rid = r.insert_record(_rec("A", abstract="text"))
        r.delete_records([rid])
        assert r.conn.execute("SELECT COUNT(*) FROM record_abstracts").fetchone()[0] == 1
    with RecordRepository(path) as r:
        assert r.conn.execute("SELECT COUNT(*) FROM record_abstracts").fetchone()[0] == 0


def test_decision_values_are_validated(repo):
    rid = repo.insert_record(_rec("A"))
    with pytest.raises(ValueError):
        repo.update_decision(rid, "ta_decision", "Inlcude")
    with pytest.raises(ValueError):
        repo.update_decision(rid, "ft_decision", "Maybe")   # not an FT option
    repo.update_decision(rid, "ta_decision", "Maybe")


def test_every_decision_change_is_audited(repo):
    rid = repo.insert_record(_rec("A"))
    repo.update_decision(rid, "ta_decision", "Include")
    repo.update_decision(rid, "ta_decision", "Include")     # no-op: not logged
    repo.update_decision(rid, "ta_decision", "Exclude")
    log = [(r["field"], r["old_value"], r["new_value"]) for r in repo.fetch_decision_log(rid)]
    assert log == [("ta_decision", "", "Include"), ("ta_decision", "Include", "Exclude")]
    assert repo.fetch_decision_log(rid)[0]["ts"].endswith("Z")


def test_duplicate_flag_changes_and_deletes_are_audited(repo):
    a = repo.insert_record(_rec("A"))
    b = repo.insert_record(_rec("B"))
    repo.set_duplicate_flag(b, "Yes", a)
    repo.unmark_duplicate(b)
    repo.delete_records([a])
    fields = [r["field"] for r in repo.fetch_decision_log()]
    assert fields.count("is_duplicate") == 2 and "__deleted__" in fields


def test_update_decisions_from_df_handles_nan_and_logs(repo):
    import pandas as pd
    rid = repo.insert_record(_rec("A"))
    df = pd.DataFrame([{"ID": rid, "TA Decision": "Include", "TA Reason": float("nan"),
                        "FT Decision": None, "FT Reason": ""}])
    repo.update_decisions_from_df(df)
    row = repo.conn.execute("SELECT ta_decision, ta_reason, ft_decision FROM records").fetchone()
    assert tuple(row) == ("Include", "", "")
    assert repo.fetch_decision_log(rid)[0]["source"] == "bulk"


def test_deleting_thousands_of_records_does_not_hit_variable_limit(repo):
    ids = repo.insert_records_bulk([(_rec("T%d" % i), "No", None) for i in range(2500)])
    repo.delete_records(ids)
    assert repo.prisma_counts()["Records identified through database searching"] == 0


def test_prisma_counts_reconcile_even_with_inconsistent_rows(repo):
    a = repo.insert_record(_rec("A"))   # TA include, FT include  -> included
    b = repo.insert_record(_rec("B"))   # TA exclude, FT include  -> invalid: ignored at FT stage
    c = repo.insert_record(_rec("C"))   # TA include, FT exclude  -> FT excluded
    d = repo.insert_record(_rec("D"))   # screened, then flagged as duplicate -> ignored everywhere
    for rid, ta, ft in ((a, "Include", "Include"), (b, "Exclude", "Include"),
                        (c, "Include", "Exclude"), (d, "Exclude", "")):
        repo.update_decision(rid, "ta_decision", ta)
        if ft:
            repo.update_decision(rid, "ft_decision", ft)
    repo.set_duplicate_flag(d, "Yes", a)
    n = repo.prisma_counts()
    assert n["Records after duplicates removed"] == 3
    assert n["Records screened (title/abstract)"] == 3
    assert n["Records excluded (title/abstract)"] == 1               # b only
    assert n["Full-text articles assessed"] == 2                     # a, c
    assert n["Full-text articles excluded (with reasons)"] == 1      # c
    assert n["Studies included in final synthesis"] == 1             # a
    assert (n["Records excluded (title/abstract)"] + n["Full-text articles assessed"]
            == n["Records screened (title/abstract)"])
    warnings = " ".join(repo.consistency_warnings())
    assert "full-text decision" in warnings and "duplicate record" in warnings


def test_database_counts_and_ft_reasons(repo):
    a = repo.insert_record(_rec("A", database="Scopus"))
    b = repo.insert_record(_rec("B", database="Scopus"))
    repo.insert_record(_rec("C", database="IEEE Xplore"))
    repo.set_duplicate_flag(b, "Yes", a)
    assert repo.database_counts() == [("IEEE Xplore", 1, 0), ("Scopus", 2, 1)]
    repo.update_decision(a, "ta_decision", "Include")
    repo.update_decision(a, "ft_decision", "Exclude")
    repo.update_decision(a, "ft_reason", "Wrong population")
    assert repo.ft_exclusion_reasons() == [("Wrong population", 1)]


# ------------------------------------------------------------- excel_io ----

def _tracker(path, with_chart=False):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw"
    ws.append(["ID", "Title", "My Notes"])
    wb.create_sheet("TA").append(["ID", "Decision", "Reviewer 2 decision"])
    wb.create_sheet("FT").append(["ID", "Final Decision"])
    if with_chart:
        from openpyxl.chart import BarChart, Reference
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=1, min_row=1, max_row=2))
        ws.add_chart(ch, "H2")
    wb.save(path)


def test_sync_keeps_user_columns_aligned_when_records_change(tmp_path, repo):
    import openpyxl
    path = str(tmp_path / "tracker.xlsx")
    _tracker(path)
    ids = [repo.insert_record(_rec(t)) for t in ("One", "Two", "Three")]
    for rid in ids:
        repo.update_decision(rid, "ta_decision", "Include")
    repo.sync_to_tracker = None  # not a Store; call the function directly
    from slr_core.excel_io import sync_to_tracker
    sync_to_tracker(repo, path, "Raw", "TA", "FT")

    # The user annotates rows by hand next to the managed columns.
    wb = openpyxl.load_workbook(path)
    for row in wb["Raw"].iter_rows(min_row=2):
        row[2].value = "note for " + str(row[1].value)
    for row in wb["TA"].iter_rows(min_row=2):
        row[2].value = "r2:" + str(row[0].value)
    wb.save(path)

    repo.delete_records([ids[0]])                       # "One" disappears
    four = repo.insert_record(_rec("Four"))             # "Four" is new
    sync_to_tracker(repo, path, "Raw", "TA", "FT")

    ws = openpyxl.load_workbook(path)["Raw"]
    rows = {r[0].value: (r[1].value, r[2].value) for r in ws.iter_rows(min_row=2) if r[0].value is not None}
    assert rows[2] == ("Two", "note for Two")           # unchanged row, note still attached
    assert rows[3] == ("Three", "note for Three")
    assert rows[four][0] == "Four"                      # appended, not squeezed into a stale row
    orphan = [r for r in ws.iter_rows(min_row=2) if r[0].value is None and r[2].value == "note for One"]
    assert len(orphan) == 1 and orphan[0][1].value is None   # managed cells cleared, note kept

    ta = openpyxl.load_workbook(path)["TA"]
    checked = 0
    for r in ta.iter_rows(min_row=2):
        if r[0].value in (2, 3):                                # the pre-existing rows
            assert r[2].value == "r2:" + str(r[0].value)       # reviewer column still lines up
            checked += 1
    assert checked == 2


def test_check_tracker_warnings(tmp_path):
    from slr_core.excel_io import check_tracker_warnings
    plain, charted = str(tmp_path / "a.xlsx"), str(tmp_path / "b.xlsx")
    _tracker(plain)
    _tracker(charted, with_chart=True)
    assert check_tracker_warnings(plain) == []
    assert any("chart" in w for w in check_tracker_warnings(charted))


def test_export_contains_prisma_detail_sheet(tmp_path, repo):
    import openpyxl
    from slr_core.excel_io import export_prisma_xlsx
    repo.insert_record(_rec("A", database="Scopus"))
    out = str(tmp_path / "r.xlsx")
    export_prisma_xlsx(repo, out)
    ws = openpyxl.load_workbook(out)["PRISMA_Detail"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Scopus" in values


# -------------------------------------------------------------- workers ----

def test_cancelled_import_delivers_no_partial_batch(tmp_path, qapp):
    import slr_workers
    f = tmp_path / "a.csv"
    f.write_text("Title\nOne\n", encoding="utf-8")
    worker = slr_workers.ImportParseWorker([str(f)], [], [], 0.9, True)
    events = []
    worker.finished_ok.connect(lambda *a: events.append("finished"))
    worker.cancelled.connect(lambda: events.append("cancelled"))
    worker.cancel()
    worker.run()                      # run synchronously on this thread
    assert events == ["cancelled"]

    worker2 = slr_workers.ImportParseWorker([str(f)], [], [], 0.9, True)
    worker2.finished_ok.connect(lambda *a: events.append("finished"))
    worker2.run()
    assert events == ["cancelled", "finished"]


# ------------------------------------------------------------------ ui ----

def test_doi_link_is_escaped_and_scheme_checked(qapp):
    from slr_ui.pages.screening import FocusView
    link = FocusView._doi_link_html
    assert link("10.1000/ABC") == '<a href="https://doi.org/10.1000/abc">10.1000/abc</a>'
    evil = link('javascript:alert(1)"><img src=x>')
    assert "<a " not in evil and "<img" not in evil
    evil2 = link('https://x.example/"><script>alert(1)</script>')
    assert "<script>" not in evil2 and "&lt;script&gt;" in evil2
    assert link("file:///C:/Windows/system32/calc.exe") == "file:///C:/Windows/system32/calc.exe"
    assert link("") == ""


def test_title_key_folds_arabic_and_persian_letter_and_digit_variants():
    # kaf: U+0643 (Arabic) vs U+06A9 (Persian); yeh: U+064A vs U+06CC
    assert title_key("\u0643\u062a\u0627\u0628") == title_key("\u06a9\u062a\u0627\u0628")
    assert title_key("\u0639\u0644\u064a") == title_key("\u0639\u0644\u06cc")
    assert title_key("\u0662\u0660\u0662\u0665") == title_key("\u06f2\u06f0\u06f2\u06f5") == title_key("2025")
    # harakat (combining marks) do not block a match
    assert title_key("\u0643\u064e\u062a\u064e\u0628\u064e") == title_key("\u06a9\u062a\u0628")
